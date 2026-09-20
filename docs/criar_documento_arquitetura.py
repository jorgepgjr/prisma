from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\Guilherme\Desktop\JavaScript_Projects\jorge\ARQUITETURA_FLUXO_TINHA.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9E2F3", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def set_font(run, name="Aptos", size=10.5, color="1F2937", bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_font(run, **kwargs)
    return run


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_font(run, size=16 if level == 1 else 12.5, color="000000", bold=True)
    return paragraph


def add_body(doc, text="", bold_lead=None):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, size=10.5, bold=True)
        add_text(p, text[len(bold_lead):], size=10.5)
    else:
        add_text(p, text, size=10.5)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.25)
    add_text(p, text, size=10.2)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    add_text(p, text, name="Cascadia Mono", size=8.5, color="334155")
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    set_cant_split(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, "1F4E79")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, value, size=9.2, color="FFFFFF", bold=True)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        for index, value in enumerate(row_data):
            cell = row.cells[index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, "F4F8FC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, str(value), size=9.1)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    add_text(p, "Arquitetura Tinhaphone, Prisma e TinhaKids  |  Documento de alinhamento", size=8, color="64748B")


def add_page_number(paragraph):
    add_text(paragraph, "Página ", size=8, color="64748B")
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("1F2937")
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Aptos Display"
        styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        styles[name].font.color.rgb = RGBColor.from_string("000000")
    styles["Body Text"].font.name = "Aptos"
    styles["Body Text"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Body Text"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")


doc = Document()
configure_styles(doc)
section = doc.sections[0]
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
add_footer(section)

# Cover and opening
title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(8)
title.paragraph_format.keep_with_next = True
add_text(title, "Arquitetura e Fluxo de Publicação Escolar", size=25, color="000000", bold=True)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
add_text(subtitle, "Do registro da foto no Tinhaphone à visualização pelos responsáveis no TinhaKids", size=13, color="365F91")

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(16)
add_text(meta, "Documento de alinhamento técnico e operacional\n", size=10, bold=True)
add_text(meta, f"Atualizado em {datetime.now().strftime('%d/%m/%Y')}", size=10, color="64748B")

add_heading(doc, "Objetivo do documento", 1)
add_body(doc, "Este documento explica, em linguagem compartilhável com o time, como funciona a cadeia completa entre os três produtos. O foco é deixar claro onde cada informação nasce, como ela é protegida e em que momento uma foto deixa de ser apenas material recebido pela escola e passa a ser uma publicação disponível para os pais.")
add_body(doc, "A decisão central da arquitetura é separar recebimento, revisão e publicação. O Tinhaphone envia a foto para o Prisma; o Prisma organiza e revisa; somente depois da seleção dos alunos e da criação explícita da publicação o TinhaKids disponibiliza o conteúdo ao responsável autorizado.")

add_heading(doc, "Resumo em uma página", 1)
add_table(doc, ["Etapa", "Sistema", "Resultado"], [
    ("1. Captura", "Tinhaphone", "Foto salva no aparelho e colocada na fila de sincronização."),
    ("2. Recebimento", "API Prisma + PostgreSQL", "Arquivo armazenado e registro Photo criado na turma."),
    ("3. Organização", "Portal Prisma", "Equipe visualiza, filtra, identifica alunos e adiciona tags."),
    ("4. Publicação", "Portal Prisma + API", "Post criado com legenda, alunos e uma ou várias fotos."),
    ("5. Distribuição", "API TinhaKids", "A publicação fica associada aos perfis familiares corretos."),
    ("6. Visualização", "Aplicativo TinhaKids", "Responsável autenticado vê apenas os filhos vinculados à conta."),
], widths=[1.05, 1.65, 4.7])

add_heading(doc, "Visão da arquitetura", 1)
add_body(doc, "Os aplicativos não conversam diretamente entre si. O Prisma funciona como o núcleo de dados e autorização: recebe a foto do Tinhaphone, oferece a operação editorial para a escola e entrega ao TinhaKids somente os dados que o responsável pode acessar.")
add_table(doc, ["Componente", "Tecnologia", "Responsabilidade"], [
    ("Tinhaphone", "Flutter / Android", "Captura, armazenamento local, autenticação da professora e sincronização resiliente."),
    ("API Prisma", "FastAPI / Python", "Upload, autorização por turma, fotos, alunos, tags, postagens, portfólios e vínculos familiares."),
    ("Banco central", "PostgreSQL", "Dados estruturados e relações entre turma, foto, aluno, criança, responsável e publicação."),
    ("Arquivos", "Pasta de uploads local", "Conteúdo original das imagens, com nome aleatório e acesso somente pela rota de mídia assinada."),
    ("Portal Prisma", "Next.js / TypeScript", "Biblioteca editorial da escola: filtros, metadados, tags, seleção, modal de publicação e vínculos."),
    ("TinhaKids", "Flutter / Riverpod", "Login familiar, seleção de criança, feed de publicações, portfólio e perfil com grid."),
], widths=[1.25, 1.55, 4.6])

add_heading(doc, "Fluxo ponta a ponta", 1)
add_heading(doc, "1. Tirar e guardar a foto", 2)
add_body(doc, "A professora tira a foto no Tinhaphone. O aplicativo grava primeiro uma cópia local no SQLite, com o estado da operação. Esse passo é importante porque a câmera não depende de a rede estar disponível no momento da captura.")
add_bullet(doc, "pending: foto aguardando tentativa de envio.")
add_bullet(doc, "syncing: envio em andamento.")
add_bullet(doc, "synced: API confirmou o recebimento e o identificador remoto foi salvo.")
add_bullet(doc, "error: tentativa falhou; a foto permanece disponível para reenvio.")

add_heading(doc, "2. Enviar para o Prisma", 2)
add_body(doc, "O Tinhaphone autentica a professora e envia a imagem por multipart para a API. A API usa o usuário do token como autor, exige a turma selecionada e verifica se aquele usuário realmente tem acesso à turma.")
add_code(doc, "POST /api/photos/upload\nAuthorization: Bearer <token>\nform-data: file, title, description, class_id")
add_body(doc, "O arquivo recebe um nome aleatório, reduzindo colisões e evitando que o nome original do aparelho seja exposto. O registro Photo guarda o caminho, turma, autor, status, data, alunos marcados e tags.")

add_heading(doc, "3. Ver a foto no Prisma", 2)
add_body(doc, "O portal consulta as fotos da turma e mostra a biblioteca em uma grade. A equipe pode filtrar por aluno, intervalo de datas e tag; alternar a organização por mês; selecionar várias imagens; abrir a foto ativa e consultar seus metadados na lateral.")
add_body(doc, "A lateral é deliberadamente informativa: data, pessoa que enviou, alunos identificados, descrição e tags. A publicação não fica permanentemente aberta nessa área; ela é criada por um modal, o que evita que o formulário ocupe o espaço principal da biblioteca.")

add_heading(doc, "4. Criar publicação ou portfólio", 2)
add_body(doc, "A escola seleciona uma ou várias fotos e abre o modal correspondente. Para uma publicação, informa a legenda e marca os alunos. Para um portfólio, informa título, descrição e objetivos pedagógicos. O backend valida que as fotos são da mesma turma e que os alunos pertencem àquela turma.")
add_code(doc, "POST /api/posts/\n{ photo_ids: [10, 11], caption: \"...\", student_ids: [2] }")
add_body(doc, "Se houver mais de uma foto, o Prisma mantém a galeria como carrossel. A mesma seleção pode gerar um projeto de portfólio para cada criança selecionada, sempre usando o perfil familiar vinculado como destino.")

add_heading(doc, "5. Entregar somente à família correta", 2)
add_body(doc, "O aluno do Prisma e a criança do TinhaKids são entidades diferentes. A relação explícita Student → Child impede que a associação seja inferida apenas pelo nome, evitando problemas com homônimos ou mudanças de grafia.")
add_table(doc, ["Cadastro escolar", "Relação", "Visão familiar"], [
    ("Student: Pedro, Grupo 2", "student_child_link", "Child: pedro"),
    ("Student: Sofia, Grupo 2", "student_child_link", "Child: sofia"),
], widths=[2.25, 1.9, 3.25])
add_body(doc, "Quando um aluno ainda não tem perfil familiar, a escola pode escolher um perfil existente ou cadastrar um novo responsável no próprio modal de vínculo. O cadastro cria a conta ativa, o perfil Child e a relação com o Student.")

add_heading(doc, "6. Login e leitura no TinhaKids", 2)
add_body(doc, "O responsável entra com e-mail ou celular. A API valida a conta, devolve um token de acesso com validade de 30 dias e lista somente as crianças vinculadas àquele responsável. O aplicativo salva o token localmente e restaura a sessão e a última criança escolhida quando é atualizado ou reaberto.")
add_code(doc, "POST /api/v1/auth/login\nGET  /api/v1/children/\nGET  /api/v1/posts/?child_id=sofia\nGET  /api/v1/portfolio/?child_id=sofia")
add_body(doc, "O backend verifica o vínculo do child_id em todas as consultas de feed e portfólio. Um responsável não consegue trocar o identificador na URL para acessar uma criança que não pertence à sua conta.")

add_heading(doc, "Como as imagens são entregues", 1)
add_body(doc, "As URLs das imagens não são caminhos públicos diretos. A API gera uma URL assinada de curta duração para cada arquivo. A rota de mídia valida a assinatura, normaliza o nome do arquivo e bloqueia tentativas de navegar para fora da pasta de uploads.")
add_body(doc, "Celulares podem produzir AVIF, HEIC ou HEIF. Esses formatos continuam guardados no original, mas a rota de mídia cria uma representação JPEG compatível com o Flutter e com diferentes navegadores. Assim, uma imagem não precisa ser reprocessada ou substituída no armazenamento para aparecer no TinhaKids.")
add_table(doc, ["Camada", "O que fica armazenado", "O que é entregue ao app"], [
    ("Original", "Arquivo AVIF, HEIC, HEIF, JPEG, PNG ou WebP", "Nunca é exposto por caminho livre."),
    ("Assinatura", "Token com caminho e expiração", "URL temporária na resposta da API."),
    ("Compatibilidade", "Conversão sob demanda quando necessário", "JPEG para renderização confiável no TinhaKids."),
], widths=[1.35, 3.0, 3.05])

add_heading(doc, "Modelo de dados essencial", 1)
add_table(doc, ["Entidade", "Representa", "Relações principais"], [
    ("Class", "Turma e ano letivo", "tem Students, Photos e professores autorizados"),
    ("Student", "Aluno no cadastro escolar", "pertence a Class; pode apontar para um Child"),
    ("Photo", "Imagem recebida pela escola", "pertence a Class; pode ter Students, Tags, Posts e Projects"),
    ("Child", "Perfil da criança no TinhaKids", "tem Parents, Posts e Projects; liga a um Student"),
    ("Parent", "Conta do responsável", "tem um ou mais Children"),
    ("Post", "Publicação no feed", "tem Children e uma galeria de Photos"),
    ("Project", "Item de portfólio", "pertence a um Child e pode ter várias Photos"),
    ("Tag", "Classificação editorial", "pode ser aplicada a várias Photos"),
], widths=[1.25, 2.25, 3.9])

add_heading(doc, "Regras de segurança e privacidade", 1)
for item in [
    "O upload usa o usuário autenticado; o cliente não escolhe arbitrariamente o autor.",
    "Professores só enxergam e usam as turmas autorizadas para aquele usuário.",
    "Publicação e portfólio só aceitam fotos da mesma turma e alunos daquela turma.",
    "A publicação só alcança crianças que possuem vínculo explícito com os alunos marcados.",
    "O feed do TinhaKids valida o responsável antes de devolver qualquer criança, postagem ou projeto.",
    "Nomes originais de arquivos não são usados como caminho de armazenamento.",
    "Em produção, a chave JWT, HTTPS, armazenamento externo e política de backup devem ser configurados fora do repositório.",
]:
    add_bullet(doc, item)

add_heading(doc, "Operação local", 1)
add_body(doc, "O ambiente local usa PostgreSQL em Docker, API FastAPI, portal Next.js e os dois aplicativos Flutter. Os atalhos PowerShell configuram o diretório temporário curto exigido pelo Gradle no Windows.")
add_table(doc, ["Serviço", "Endereço local", "Atalho"], [
    ("API Prisma", "http://127.0.0.1:8000", "iniciar-prisma.ps1"),
    ("Portal Prisma", "http://localhost:3000", "iniciar-prisma.ps1"),
    ("TinhaKids web", "Flutter em Chrome", "iniciar-tinhakids.ps1"),
    ("Tinhaphone", "Emulador Android", "iniciar-tinhaphone.ps1"),
], widths=[1.45, 2.05, 3.9])
add_body(doc, "Para recriar o banco de demonstração, o script de seed limpa as tabelas e carrega turmas, alunos, famílias, fotos e dados de exemplo. Essa operação é destrutiva para o banco local e não deve ser usada em uma base de produção.")

add_heading(doc, "Demonstração recomendada", 1)
for item in [
    "Iniciar o Prisma e confirmar que a API e o portal estão disponíveis.",
    "Abrir o Tinhaphone, autenticar a professora e tirar uma foto no Grupo 2.",
    "Aguardar a confirmação de sincronização ou usar o reenvio se a rede estiver indisponível.",
    "Abrir a turma no Prisma, localizar a foto recebida e conferir data, autor e arquivo.",
    "Selecionar a foto, abrir Publicar, escrever uma legenda e marcar Pedro ou Sofia.",
    "Entrar no TinhaKids com a conta familiar de demonstração e escolher a criança marcada.",
    "Conferir a publicação no feed e, quando aplicável, no Perfil e no grid de fotos.",
]:
    add_bullet(doc, item)

add_heading(doc, "Estado atual e próximos cuidados", 1)
add_body(doc, "O fluxo completo está implementado e validado localmente: upload, armazenamento, visualização no Prisma, vínculo, publicação, leitura protegida e compatibilidade de imagem. O próximo passo para produção é trocar as credenciais de demonstração, configurar segredo JWT fora do código, usar HTTPS, mover os arquivos para armazenamento com backup e definir a política de retenção e consentimento de imagens.")
add_body(doc, "Para desenvolvimento, os testes automatizados cobrem publicação, carrossel, portfólio, isolamento entre responsáveis e cadastro de família. Também foram validados o upload AVIF, a conversão de mídia para JPEG e o build dos aplicativos.")

# Page-number footer after all content
for sec in doc.sections:
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.clear()
    add_page_number(p)

doc.core_properties.title = "Arquitetura e Fluxo de Publicação Escolar"
doc.core_properties.subject = "Tinhaphone, Prisma e TinhaKids"
doc.core_properties.author = "Equipe do projeto"
doc.core_properties.comments = "Documento de alinhamento técnico e operacional"
doc.save(OUT)
print(OUT)
